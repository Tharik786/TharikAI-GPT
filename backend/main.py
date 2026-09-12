import os
import sys
import json
from pathlib import Path
from dotenv import load_dotenv

# Fix Windows console encoding if needed
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Robustly load .env and sys.path from backend directory regardless of working directory
current_path = Path(__file__).resolve()
backend_dir = current_path.parent if current_path.parent.name == "backend" else current_path.parent / "backend"
root_dir = backend_dir.parent

if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

# Change directory to backend so uvicorn app resolution works cleanly
try:
    os.chdir(str(backend_dir))
except Exception:
    pass

load_dotenv(backend_dir / ".env")
load_dotenv()


import io
import time
import socket
import shutil
import subprocess
import urllib.parse
from contextlib import asynccontextmanager

import httpx
from fastapi import FastAPI, HTTPException, UploadFile, File, Response, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from app.llmModel import stream_chat_completion, GeminiError
from app.webService import should_perform_web_search, perform_web_search
from app.imageGenerator import (
    detect_image_prompt,
    generate_image_bytes,
    generate_kie_image,
)
from app.docGenerator import (
    render_carbone_document,
    fetch_rendered_file,
    detect_document_prompt,
    strip_template_metadata,
)
from app.db import (
    init_db,
    get_db_connection,
    get_user,
    save_user,
    list_conversations,
    save_conversation,
    set_messages as db_set_messages,
    delete_conversation as db_delete_conversation,
)

# ============================================================
# Frontend Auto-Run & Integration Configuration
# ============================================================
SERVE_FRONTEND = os.getenv("SERVE_FRONTEND", "true").strip().lower() in ("true", "1", "yes")
AUTO_RUN_FRONTEND = os.getenv("AUTO_RUN_FRONTEND", "true").strip().lower() in ("true", "1", "yes")
FRONTEND_DIST_DIR_ENV = os.getenv("FRONTEND_DIST_DIR", "").strip()
VITE_DEV_HOST = os.getenv("VITE_DEV_HOST", "127.0.0.1")
VITE_DEV_PORT = int(os.getenv("VITE_DEV_PORT", "5173"))
VITE_DEV_URL = f"http://{VITE_DEV_HOST}:{VITE_DEV_PORT}"


def get_frontend_dir() -> Path | None:
    candidates = [
        backend_dir.parent / "frontend",
        backend_dir / "frontend",
        Path.cwd() / "frontend",
        Path.cwd(),
    ]
    for p in candidates:
        if p.exists() and (p / "package.json").exists():
            return p.resolve()
    return None


def get_frontend_dist_dir() -> Path | None:
    if FRONTEND_DIST_DIR_ENV:
        p = Path(FRONTEND_DIST_DIR_ENV).resolve()
        if p.exists():
            return p
    candidates = [
        backend_dir.parent / "frontend" / "dist",
        backend_dir / "frontend" / "dist",
        backend_dir / "dist",
        Path.cwd() / "frontend" / "dist",
        Path.cwd() / "dist",
    ]
    for p in candidates:
        if p.exists() and (p / "index.html").exists():
            return p.resolve()
    return candidates[0].resolve()


def is_vite_dev_server_running() -> bool:
    try:
        with socket.create_connection((VITE_DEV_HOST, VITE_DEV_PORT), timeout=0.3):
            return True
    except (OSError, ConnectionRefusedError):
        return False


@asynccontextmanager
async def lifespan(app: FastAPI):
    # 1. Initialize cloud PostgreSQL tables on startup
    init_db()

    # 2. Automatically run/build frontend when backend starts
    frontend_process = None
    if AUTO_RUN_FRONTEND:
        # Avoid double-spawning if Vite dev server is already running
        if not is_vite_dev_server_running():
            frontend_path = get_frontend_dir()
            if frontend_path and shutil.which("npm"):
                npm_cmd = "npm.cmd" if sys.platform == "win32" else "npm"
                dist_dir = get_frontend_dist_dir()
                dist_ready = dist_dir and dist_dir.exists() and (dist_dir / "index.html").exists()
                is_prod = bool(os.getenv("RENDER") or os.getenv("PRODUCTION"))

                if not is_prod:
                    try:
                        print(f"[Auto-Runner] Starting frontend dev server in {frontend_path}...")
                        frontend_process = subprocess.Popen(
                            [npm_cmd, "run", "dev", "--", "--host", VITE_DEV_HOST, "--port", str(VITE_DEV_PORT)],
                            cwd=str(frontend_path),
                            stdout=subprocess.DEVNULL,
                            stderr=subprocess.DEVNULL,
                            shell=(sys.platform == "win32"),
                        )
                        print(f"[Auto-Runner] Frontend dev server started (PID: {frontend_process.pid}) at {VITE_DEV_URL}")
                    except Exception as e:
                        print(f"[Auto-Runner] Could not start frontend dev server: {e}")
                elif not dist_ready:
                    try:
                        print(f"[Auto-Runner] Building frontend in {frontend_path}...")
                        subprocess.run([npm_cmd, "run", "build"], cwd=str(frontend_path), check=True, shell=(sys.platform == "win32"))
                        print("[Auto-Runner] Frontend build completed.")
                    except Exception as e:
                        print(f"[Auto-Runner] Frontend build failed: {e}")

    yield

    # Clean up frontend process on backend shutdown
    if frontend_process:
        print("[Auto-Runner] Stopping frontend process...")
        try:
            if sys.platform == "win32":
                subprocess.call(["taskkill", "/F", "/T", "/PID", str(frontend_process.pid)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            else:
                frontend_process.terminate()
                frontend_process.wait(timeout=3)
        except Exception:
            try:
                frontend_process.kill()
            except Exception:
                pass


app = FastAPI(title="TharikAI API", lifespan=lifespan)

cors_env = os.getenv("CORS_ORIGINS", "*").strip()
if cors_env == "*":
    app.add_middleware(
        CORSMiddleware,
        allow_origin_regex=".*",
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )
else:
    origins = [o.strip() for o in cors_env.split(",") if o.strip()]
    app.add_middleware(
        CORSMiddleware,
        allow_origins=origins,
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )

# Mount Vite static assets if dist exists
_initial_dist = get_frontend_dist_dir()
if SERVE_FRONTEND and _initial_dist and _initial_dist.exists():
    _assets_dir = _initial_dist / "assets"
    if _assets_dir.exists() and _assets_dir.is_dir():
        app.mount("/assets", StaticFiles(directory=str(_assets_dir)), name="frontend_assets")


async def render_frontend_or_spa(full_path: str, request: Request):
    clean_path = full_path.lstrip("/")
    # Exclude API endpoints, Swagger docs, and OpenAPI schema from rewrites
    if clean_path.startswith("api/") or clean_path == "api" or clean_path in ("docs", "redoc", "openapi.json"):
        raise HTTPException(status_code=404, detail="Not Found")

    # 1. First priority: Check if static dist directory exists (pre-bundled production assets)
    dist_root = get_frontend_dist_dir()
    if dist_root and dist_root.exists():
        if clean_path:
            requested_file = (dist_root / clean_path).resolve()
            if str(requested_file).startswith(str(dist_root.resolve())) and requested_file.is_file():
                return FileResponse(requested_file)

        # SPA fallback rewrite: return dist/index.html for root or client-side routes
        index_file = dist_root / "index.html"
        if index_file.is_file():
            return FileResponse(index_file)

    # 2. Second priority: If dist is not yet built, proxy to Vite dev server on port 5173
    if is_vite_dev_server_running():
        target_url = f"{VITE_DEV_URL}/{clean_path}"
        if request.url.query:
            target_url += f"?{request.url.query}"
        try:
            async with httpx.AsyncClient() as client:
                resp = await client.get(target_url, timeout=10.0, follow_redirects=True)
                content_type = resp.headers.get("content-type", "text/html")
                headers = {
                    k: v
                    for k, v in resp.headers.items()
                    if k.lower() not in ("content-length", "content-encoding", "transfer-encoding", "connection")
                }
                return Response(
                    content=resp.content,
                    status_code=resp.status_code,
                    media_type=content_type,
                    headers=headers,
                )
        except Exception:
            pass

    return HTMLResponse(
        "<!doctype html><html><head><title>TharikAI</title></head>"
        "<body style='font-family:sans-serif;display:flex;align-items:center;justify-content:center;height:100vh;margin:0;background:#0f0f10;color:#f0f0f0;'>"
        "<div style='text-align:center;max-width:500px;padding:2rem;border:1px solid #27272a;border-radius:12px;background:#18181b;'>"
        "<h2 style='margin-bottom:0.5rem;'>[+] TharikAI Backend Ready</h2>"
        "<p style='color:#a1a1aa;'>Frontend is starting up or building...</p>"
        "<p style='font-size:0.875rem;color:#71717a;'>Once ready, refresh this page or visit port 5173.</p>"
        "</div></body></html>",
        status_code=200,
    )


class ChatMessage(BaseModel):
    role: str  # "user" | "assistant"
    content: str
    images: list[dict | str] | None = None


class ChatBody(BaseModel):
    messages: list[ChatMessage]
    email: str | None = None


class RegisterBody(BaseModel):
    email: str
    name: str | None = None
    password_hash: str


class LoginBody(BaseModel):
    email: str
    password_hash: str


class ConversationBody(BaseModel):
    id: str
    email: str
    title: str | None = None
    createdAt: int | None = None
    updatedAt: int | None = None


class MessageItem(BaseModel):
    id: str | None = None
    role: str
    content: str
    createdAt: int | None = None
    images: list[dict | str] | None = None


class MessagesBody(BaseModel):
    id: str
    messages: list[MessageItem]
    updatedAt: int | None = None


class ImageBody(BaseModel):
    prompt: str
    aspect_ratio: str | None = "auto"
    resolution: str | None = "1K"
    background: str | None = "auto"


@app.post("/api/image")
async def image_endpoint(body: ImageBody):
    """
    Direct AI image generation endpoint proxying to GPT Image 2.5 Flare (Kie.ai).
    Returns RAW image binary bytes (PNG/JPEG) for frontend <img> and Blob display.
    """
    if not body.prompt or not body.prompt.strip():
        raise HTTPException(status_code=400, detail="Prompt cannot be empty")

    raw_bytes, content_type = await generate_image_bytes(
        prompt=body.prompt.strip(),
        aspect_ratio=body.aspect_ratio or "auto",
        resolution=body.resolution or "1K",
        background=body.background or "auto",
    )
    return Response(
        content=raw_bytes,
        media_type=content_type,
        headers={
            "Cache-Control": "public, max-age=86400",
            "Content-Type": content_type,
            "Content-Disposition": f'inline; filename="image_{int(time.time())}.png"',
        },
    )


@app.get("/api/image")
async def image_get_endpoint(
    prompt: str,
    aspect_ratio: str = "auto",
    resolution: str = "1K",
    background: str = "auto",
):
    """
    Direct AI image generation GET endpoint returning RAW image binary data.
    Allows markdown ![alt](/api/image?prompt=...) to directly render in browsers.
    """
    if not prompt or not prompt.strip():
        raise HTTPException(status_code=400, detail="Prompt cannot be empty")

    raw_bytes, content_type = await generate_image_bytes(
        prompt=prompt.strip(),
        aspect_ratio=aspect_ratio,
        resolution=resolution,
        background=background,
    )
    return Response(
        content=raw_bytes,
        media_type=content_type,
        headers={
            "Cache-Control": "public, max-age=86400",
            "Content-Type": content_type,
            "Content-Disposition": f'inline; filename="image_{int(time.time())}.png"',
        },
    )


@app.post("/api/generate-image")
async def generate_image_json_endpoint(body: ImageBody):
    """
    Returns structured JSON with the public image URL, task ID, prompt, and provider metadata.
    """
    if not body.prompt or not body.prompt.strip():
        raise HTTPException(status_code=400, detail="Prompt cannot be empty")

    return await generate_kie_image(
        prompt=body.prompt.strip(),
        aspect_ratio=body.aspect_ratio or "auto",
        resolution=body.resolution or "1K",
        background=body.background or "auto",
    )


# ============================================================
# Carbone.io Document Generator Endpoints
# ============================================================

class DocumentGenerateBody(BaseModel):
    title: str | None = None
    content: str | None = None
    query: str | None = None
    use_search: bool | None = False


@app.post("/api/documents/generate")
async def generate_document_endpoint(body: DocumentGenerateBody):
    """
    Generate an executive-styled PDF document via Carbone.io API v4.
    Accepts direct markdown content OR a query/prompt to synthesize with LLM + search.
    """
    title = (body.title or "").strip()
    content = (body.content or "").strip()
    query = (body.query or "").strip()

    if not content and not query:
        raise HTTPException(status_code=400, detail="Either 'content' or 'query' must be provided")

    # If query is provided without content, synthesize with LLM
    if not content and query:
        if not title:
            title = query[:60].strip()

        search_context = ""
        if body.use_search or should_perform_web_search(query):
            search_context = await perform_web_search(query)

        llm_messages = [
            {
                "role": "user",
                "content": (
                    f"Write an in-depth, comprehensive analysis on the topic: '{query}'.\n\n"
                    "Include an Executive Summary, Core Analysis with detailed insights, "
                    "a Comparative Table if helpful, and Strategic Next Steps. Format in clean markdown with clear headings.\n\n"
                    "CRITICAL: Do NOT output any template metadata, header blocks, 'Executive Document:', 'Date:', 'Prepared by:', 'Author:', or horizontal rule dividers at the top. Begin directly with the first section heading."
                ),
            }
        ]

        text_chunks = []
        try:
            async for chunk in stream_chat_completion(llm_messages, web_search_context=search_context):
                text_chunks.append(chunk)
            content = strip_template_metadata("".join(text_chunks))
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"LLM content generation failed: {str(e)}")

    if not title:
        title = "AI Document"

    try:
        render_res = render_carbone_document(title=title, markdown_content=content)
        render_id = render_res["renderId"]
        filename = render_res["filename"]
        encoded_fn = urllib.parse.quote(filename)

        return {
            "success": True,
            "renderId": render_id,
            "title": title,
            "filename": filename,
            "downloadUrl": f"/api/documents/download/{render_id}?filename={encoded_fn}",
            "viewUrl": f"/api/documents/view/{render_id}?filename={encoded_fn}",
            "createdAt": render_res.get("createdAt"),
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Carbone document generation error: {str(e)}")


@app.get("/api/documents/download/{render_id:path}")
async def download_document_endpoint(render_id: str, filename: str = "document.pdf"):
    """
    Downloads the rendered document bytes from Carbone with Content-Disposition: attachment.
    """
    try:
        pdf_bytes = fetch_rendered_file(render_id)
        clean_filename = urllib.parse.quote(filename)
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{clean_filename}',
                "Content-Type": "application/pdf",
                "Cache-Control": "public, max-age=86400",
            },
        )
    except Exception as e:
        raise HTTPException(status_code=404, detail=f"Document download failed: {str(e)}")


@app.get("/api/documents/view/{render_id:path}")
async def view_document_endpoint(render_id: str, filename: str = "document.pdf"):
    """
    Serves rendered PDF inline for browser viewing / iframe preview.
    """
    try:
        pdf_bytes = fetch_rendered_file(render_id)
        clean_filename = urllib.parse.quote(filename)
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'inline; filename="{filename}"; filename*=UTF-8\'\'{clean_filename}',
                "Content-Type": "application/pdf",
                "Cache-Control": "public, max-age=86400",
            },
        )
    except Exception as e:
        raise HTTPException(status_code=404, detail=f"Document view failed: {str(e)}")


@app.api_route("/", methods=["GET", "HEAD"])
async def root(request: Request):
    if SERVE_FRONTEND:
        return await render_frontend_or_spa("", request)
    return {"status": "ok", "message": "TharikAI API is running", "docs": "/docs"}


@app.api_route("/api/health", methods=["GET", "HEAD"])
async def health():

    db_status = "disconnected"
    try:
        con = get_db_connection()
        con.run("SELECT 1")
        con.close()
        db_status = "connected"
    except Exception as e:
        db_status = f"error: {str(e)}"

    return {
        "status": "ok" if db_status == "connected" else "degraded",
        "database": db_status,
    }






@app.post("/api/auth/register")
async def register(body: RegisterBody):
    email = body.email.strip().lower()
    existing = get_user(email)
    if existing:
        raise HTTPException(status_code=400, detail="An account with this email already exists.")
    user = save_user(email, body.name or email.split("@")[0], body.password_hash)
    return {"success": True, "user": {"email": user["email"], "name": user["name"]}}


@app.post("/api/auth/login")
async def login(body: LoginBody):
    email = body.email.strip().lower()
    user = get_user(email)
    if not user or user.get("password_hash") != body.password_hash:
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    return {"success": True, "user": {"email": user["email"], "name": user["name"]}}


@app.get("/api/conversations")
async def get_user_conversations(email: str):
    email = email.strip().lower()
    convs = list_conversations(email)
    return {"conversations": convs}


@app.post("/api/conversations")
async def create_or_update_conversation(body: ConversationBody):
    save_conversation(body.id, body.email, body.title, body.createdAt, body.updatedAt)
    return {"success": True}


@app.post("/api/conversations/messages")
async def update_conversation_messages(body: MessagesBody):
    db_set_messages(body.id, body.messages, body.updatedAt)
    return {"success": True}


@app.delete("/api/conversations/{conv_id}")
async def remove_conversation(conv_id: str):
    db_delete_conversation(conv_id)
    return {"success": True}


@app.post("/api/extract-document")
async def extract_document(file: UploadFile = File(...)):
    """
    Extracts clean readable text from uploaded documents (PDF, DOCX, TXT, MD, CSV, code).
    Prevents binary byte corruption and enables AI models to read and analyze files.
    """
    filename = file.filename or "uploaded_file"
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    content = await file.read()
    
    extracted_text = ""
    page_count = 1

    if ext == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            page_count = len(reader.pages)
            pages_text = []
            for idx, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                if txt.strip():
                    pages_text.append(f"--- Page {idx + 1} ---\n{txt.strip()}")
            extracted_text = "\n\n".join(pages_text).strip()
            if not extracted_text:
                extracted_text = "[Notice: This PDF contains no extractable text. It may contain scanned images or protected content.]"
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")

    elif ext in ("docx", "doc"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            extracted_text = "\n\n".join(paragraphs).strip()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read Word document: {str(e)}")

    else:
        # Plain text, Markdown, CSV, JSON, code files
        try:
            extracted_text = content.decode("utf-8")
        except UnicodeDecodeError:
            extracted_text = content.decode("latin-1", errors="ignore")

    return {
        "success": True,
        "filename": filename,
        "page_count": page_count,
        "size": len(content),
        "text": extracted_text,
    }


@app.post("/api/chat")
async def chat(body: ChatBody):
    """
    Streaming chat completion endpoint with multimodal vision and real-time internet search grounding.
    """
    if not body.messages:
        raise HTTPException(status_code=400, detail="messages cannot be empty")

    llm_messages = []
    for m in body.messages:
        msg_dict = {"role": m.role, "content": m.content}
        if m.images:
            msg_dict["images"] = m.images
        llm_messages.append(msg_dict)

    # Detect if user query is a question/search topic requiring fresh internet search results
    latest_user_text = ""
    for m in reversed(llm_messages):
        if m.get("role") == "user":
            latest_user_text = m.get("content", "")
            break

    # Detect if user query is requesting image generation
    image_prompt = detect_image_prompt(latest_user_text)
    if image_prompt:
        async def image_event_stream():
            yield f"data: {json.dumps({'type': 'search_status', 'status': f'🎨 Creating image with GPT Image 2.5 Flare for \"{image_prompt[:45]}\"...'})}\n\n"
            encoded_prompt = urllib.parse.quote(image_prompt)
            img_markdown = f"![{image_prompt}](/api/image?prompt={encoded_prompt})\n\n"
            yield f"data: {json.dumps({'delta': img_markdown})}\n\n"
            yield f"data: {json.dumps({'done': True})}\n\n"

        return StreamingResponse(image_event_stream(), media_type="text/event-stream")

    # Detect if user query is requesting document generation
    doc_topic = detect_document_prompt(latest_user_text)
    if doc_topic:
        async def doc_event_stream():
            yield f"data: {json.dumps({'type': 'search_status', 'status': f'📄 Generating document for \"{doc_topic[:45]}\"...'})}\n\n"

            search_ctx = ""
            if should_perform_web_search(doc_topic):
                search_ctx = await perform_web_search(doc_topic)

            doc_title = doc_topic.title()
            llm_messages_doc = [
                {
                    "role": "user",
                    "content": (
                        f"Write an in-depth, comprehensive, well-structured analysis on: '{doc_topic}'.\n\n"
                        "Include an Executive Summary, Key Findings, Detailed Sections with relevant data, "
                        "and Strategic Recommendations. Format in clean markdown with clear headings.\n\n"
                        "CRITICAL: Do NOT include any template headers, document titles, dates, author names, 'Executive Document:', "
                        "'Prepared by:', 'Date:', metadata fields, or horizontal lines at the beginning. Start immediately with the first section heading."
                    ),
                }
            ]

            full_chunks = []
            try:
                async for chunk in stream_chat_completion(llm_messages_doc, web_search_context=search_ctx):
                    full_chunks.append(chunk)
                    yield f"data: {json.dumps({'delta': chunk})}\n\n"
                doc_content = strip_template_metadata("".join(full_chunks))
            except Exception as e:
                yield f"data: {json.dumps({'error': f'Failed to generate document content: {str(e)}'})}\n\n"
                return

            try:
                yield f"data: {json.dumps({'type': 'search_status', 'status': f'📄 Rendering PDF document...'})}\n\n"
                render_res = render_carbone_document(title=doc_title, markdown_content=doc_content)
                render_id = render_res["renderId"]
                filename = render_res["filename"]
                encoded_fn = urllib.parse.quote(filename)

                # Stream clean download link at the end of the text
                doc_link = f"\n\n[📄 Download {doc_title}.pdf](/api/documents/download/{render_id}?filename={encoded_fn})\n\n"
                yield f"data: {json.dumps({'delta': doc_link})}\n\n"
            except Exception as e:
                yield f"data: {json.dumps({'error': f'Carbone document generation error: {str(e)}'})}\n\n"
                return

            yield f"data: {json.dumps({'done': True})}\n\n"

        return StreamingResponse(doc_event_stream(), media_type="text/event-stream")

    search_context = ""
    if should_perform_web_search(latest_user_text):
        search_context = await perform_web_search(latest_user_text)

    async def event_stream():
        try:
            async for chunk in stream_chat_completion(llm_messages, web_search_context=search_context):
                yield f"data: {json.dumps({'delta': chunk})}\n\n"
        except GeminiError as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"
            return
        yield f"data: {json.dumps({'done': True})}\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


# ============================================================
# Unified Frontend Reverse Proxy & SPA Catch-All Rewrite
# ============================================================
if SERVE_FRONTEND:
    @app.api_route("/{full_path:path}", methods=["GET", "HEAD"], include_in_schema=False)
    async def serve_spa_route(full_path: str, request: Request):
        return await render_frontend_or_spa(full_path, request)


if __name__ == "__main__":
    import uvicorn

    port = int(os.getenv("PORT", "8000"))
    host = os.getenv("HOST", "0.0.0.0")
    print("=" * 60)
    print(">>> Launching TharikAI (Backend + Frontend)")
    print(f">>> Serving at: http://localhost:{port}")
    print("=" * 60)
    uvicorn.run("main:app", host=host, port=port, reload=True)

